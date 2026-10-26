# -*- coding: utf-8 -*-
"""
Sistema de RRHH LOGITEx — versión web para Vercel.

Arquitectura
  index.html          interfaz (sin datos personales)
  app.py              este archivo: la API (Flask) y la entrega de index.html
  Base de datos       PostgreSQL en la nube (Neon, desde el Marketplace de Vercel)

Variables de entorno (Vercel → Settings → Environment Variables)
  DATABASE_URL        la crea sola la integración de Neon (también se acepta POSTGRES_URL)
  ADMIN_USERNAME      usuario administrador inicial (solo se usa si no hay ningún usuario)
  ADMIN_PASSWORD      su contraseña (mínimo 8 caracteres)
  ADMIN_RESET_PASSWORD  opcional: restablece la contraseña de ADMIN_USERNAME (se aplica una vez)

Sin DATABASE_URL usa un SQLite local (solo para pruebas en tu PC):
  pip install -r requirements.txt
  python app.py              → http://127.0.0.1:5000
"""

import csv
import datetime
import hashlib
import hmac
import io
import json
import os
import secrets
import sqlite3
import ssl
import time
import urllib.parse

from flask import Flask, Response, jsonify, request, send_file

try:
    import pg8000.dbapi as pg
except ImportError:  # pragma: no cover
    pg = None

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    XLSX = True
except ImportError:  # pragma: no cover
    XLSX = False

APP_VERSION = "12.0-web"
DURACION_SESION = 12 * 3600
MAX_CUERPO = 4 * 1024 * 1024          # Vercel acepta hasta 4,5 MB por pedido
RESPALDOS_A_CONSERVAR = 30
INTENTOS_LOGIN = 8                    # por usuario cada 15 minutos
COOKIE = "logitex_sesion"
COLECCIONES = ["vacPeriodos", "evalData", "empEquiposData", "fotos", "tabPlantas",
               "tabSectores", "tabPuestos", "tabActividades", "inventarioITAplicado"]

DIR_RAIZ = os.path.dirname(os.path.abspath(__file__))
def _buscar_url_db():
    """Neon/Vercel pueden crear la variable con distintos nombres (o con un prefijo)."""
    for nombre in ("DATABASE_URL", "POSTGRES_URL", "NEON_DATABASE_URL", "POSTGRES_PRISMA_URL", "DATABASE_URL_UNPOOLED"):
        if os.environ.get(nombre, "").startswith(("postgres://", "postgresql://")):
            return nombre, os.environ[nombre]
    for nombre, valor in sorted(os.environ.items()):
        if nombre.endswith(("DATABASE_URL", "POSTGRES_URL")) and valor.startswith(("postgres://", "postgresql://")):
            return nombre, valor
    return "", ""


DB_VAR, DB_URL = _buscar_url_db()
EN_VERCEL = bool(os.environ.get("VERCEL"))
ES_PG = DB_URL.startswith(("postgres://", "postgresql://"))

app = Flask(__name__, static_folder=None)
app.config["MAX_CONTENT_LENGTH"] = MAX_CUERPO


class ErrorAPI(Exception):
    def __init__(self, codigo, mensaje):
        super().__init__(mensaje)
        self.codigo = codigo


def ahora_iso():
    return datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds")


# =============================================================================
# BASE DE DATOS (PostgreSQL en Vercel, SQLite para pruebas locales)
# =============================================================================

def _conectar_pg():
    u = urllib.parse.urlparse(DB_URL)
    params = urllib.parse.parse_qs(u.query)
    local = u.hostname in ("localhost", "127.0.0.1")
    usar_ssl = params.get("sslmode", ["disable" if local else "require"])[0] != "disable"
    return pg.connect(
        user=urllib.parse.unquote(u.username or ""),
        password=urllib.parse.unquote(u.password or ""),
        host=u.hostname, port=u.port or 5432,
        database=(u.path or "/").lstrip("/") or "postgres",
        ssl_context=ssl.create_default_context() if usar_ssl else None,
        timeout=15,
    )


class Db:
    """Conexión por pedido. Usa '?' como marcador en ambos motores."""

    def __init__(self):
        if ES_PG:
            self.con = _conectar_pg()
        else:
            self.con = sqlite3.connect(os.environ.get("SQLITE_PATH", os.path.join(DIR_RAIZ, "rrhh_local.db")), timeout=15)

    def ejecutar(self, sql, params=()):
        cur = self.con.cursor()
        cur.execute(sql.replace("?", "%s") if ES_PG else sql, tuple(params))
        return cur

    def filas(self, sql, params=()):
        cur = self.ejecutar(sql, params)
        if cur.description is None:
            return []
        cols = [d[0] for d in cur.description]
        return [dict(zip(cols, r)) for r in cur.fetchall()]

    def fila(self, sql, params=()):
        r = self.filas(sql, params)
        return r[0] if r else None

    def __enter__(self):
        return self

    def __exit__(self, tipo, *_):
        try:
            if tipo is None:
                self.con.commit()
            else:
                self.con.rollback()
        finally:
            self.con.close()


ESQUEMA = [
    "CREATE TABLE IF NOT EXISTS meta (clave TEXT PRIMARY KEY, valor TEXT)",
    """CREATE TABLE IF NOT EXISTS usuarios (
        id TEXT PRIMARY KEY, username TEXT NOT NULL UNIQUE, nombre TEXT NOT NULL,
        rol TEXT NOT NULL, activo INTEGER NOT NULL DEFAULT 1, pass_hash TEXT NOT NULL,
        creado TEXT, ultimo_acceso TEXT)""",
    """CREATE TABLE IF NOT EXISTS empleados (
        id TEXT PRIMARY KEY, legajo TEXT, nombre TEXT, planta TEXT, sector TEXT,
        activo INTEGER, datos TEXT NOT NULL, actualizado TEXT)""",
    "CREATE TABLE IF NOT EXISTS colecciones (clave TEXT PRIMARY KEY, datos TEXT NOT NULL, actualizado TEXT)",
    "CREATE TABLE IF NOT EXISTS sesiones (token_hash TEXT PRIMARY KEY, usuario_id TEXT NOT NULL, expira DOUBLE PRECISION NOT NULL)",
    ("CREATE TABLE IF NOT EXISTS historial (id SERIAL PRIMARY KEY, fecha TEXT NOT NULL, usuario TEXT, accion TEXT NOT NULL, detalle TEXT)"
     if ES_PG else
     "CREATE TABLE IF NOT EXISTS historial (id INTEGER PRIMARY KEY AUTOINCREMENT, fecha TEXT NOT NULL, usuario TEXT, accion TEXT NOT NULL, detalle TEXT)"),
    ("CREATE TABLE IF NOT EXISTS respaldos (id SERIAL PRIMARY KEY, fecha TEXT NOT NULL, usuario TEXT, motivo TEXT, empleados INTEGER, datos TEXT NOT NULL)"
     if ES_PG else
     "CREATE TABLE IF NOT EXISTS respaldos (id INTEGER PRIMARY KEY AUTOINCREMENT, fecha TEXT NOT NULL, usuario TEXT, motivo TEXT, empleados INTEGER, datos TEXT NOT NULL)"),
    "INSERT INTO meta (clave, valor) VALUES ('version_datos', '0') ON CONFLICT (clave) DO NOTHING",
]

_inicializada = False


def inicializar():
    """Crea las tablas (una vez por arranque) y el administrador inicial si hace falta."""
    global _inicializada
    if _inicializada:
        return
    with Db() as db:
        for sql in ESQUEMA:
            db.ejecutar(sql)
        hay = db.fila("SELECT COUNT(*) AS n FROM usuarios")["n"]
        user = (os.environ.get("ADMIN_USERNAME") or "").strip().lower()
        clave = os.environ.get("ADMIN_PASSWORD") or ""
        if not hay and user and len(clave) >= 8:
            db.ejecutar("INSERT INTO usuarios (id, username, nombre, rol, activo, pass_hash, creado) VALUES (?,?,?,?,?,?,?)",
                        ("u" + secrets.token_hex(6), user, user, "admin", 1, hash_password(clave), ahora_iso()))
            registrar(db, "sistema", "crear_usuario", user + " (administrador inicial)")
        restablecer_admin(db, user)
        db.ejecutar("DELETE FROM sesiones WHERE expira < ?", (time.time(),))
    _inicializada = True


def restablecer_admin(db, user):
    """Restablece la contraseña del administrador con la variable ADMIN_RESET_PASSWORD.

    Se aplica una sola vez por cada valor nuevo de la variable: si después cambiás la
    contraseña desde el sistema, no se vuelve a pisar en el próximo arranque.
    """
    clave = os.environ.get("ADMIN_RESET_PASSWORD") or ""
    if not user or len(clave) < 8:
        return
    marca = hashlib.sha256(f"{user}:{clave}".encode()).hexdigest()
    previa = db.fila("SELECT valor FROM meta WHERE clave = 'admin_reset'")
    if previa and previa["valor"] == marca:
        return
    r = db.fila("SELECT id FROM usuarios WHERE username = ?", (user,))
    if r:
        db.ejecutar("UPDATE usuarios SET pass_hash = ?, activo = 1, rol = 'admin' WHERE id = ?", (hash_password(clave), r["id"]))
        db.ejecutar("DELETE FROM sesiones WHERE usuario_id = ?", (r["id"],))
    else:
        db.ejecutar("INSERT INTO usuarios (id, username, nombre, rol, activo, pass_hash, creado) VALUES (?,?,?,?,?,?,?)",
                    ("u" + secrets.token_hex(6), user, user, "admin", 1, hash_password(clave), ahora_iso()))
    # Levanta el bloqueo por intentos fallidos de ese usuario
    db.ejecutar("DELETE FROM historial WHERE accion = 'login_fallido' AND usuario = ?", (user,))
    db.ejecutar("INSERT INTO meta (clave, valor) VALUES ('admin_reset', ?) "
                "ON CONFLICT (clave) DO UPDATE SET valor = excluded.valor", (marca,))
    registrar(db, "sistema", "reset_password", f"{user} (variable ADMIN_RESET_PASSWORD)")


def registrar(db, usuario, accion, detalle=""):
    db.ejecutar("INSERT INTO historial (fecha, usuario, accion, detalle) VALUES (?,?,?,?)",
                (ahora_iso(), (usuario or "")[:60], accion, (detalle or "")[:300]))


# =============================================================================
# CONTRASEÑAS Y SESIONES
# =============================================================================

_ITER = 240_000


def hash_password(clave):
    sal = secrets.token_bytes(16)
    dk = hashlib.pbkdf2_hmac("sha256", clave.encode(), sal, _ITER)
    return f"pbkdf2_sha256${_ITER}${sal.hex()}${dk.hex()}"


def verificar_password(clave, guardado):
    try:
        alg, it, sal, h = guardado.split("$")
        dk = hashlib.pbkdf2_hmac("sha256", clave.encode(), bytes.fromhex(sal), int(it))
        return alg == "pbkdf2_sha256" and hmac.compare_digest(dk.hex(), h)
    except (ValueError, AttributeError):
        return False


def _hash_token(token):
    return hashlib.sha256(token.encode()).hexdigest()


def _usuario_publico(r):
    return {"id": r["id"], "username": r["username"], "nombre": r["nombre"], "rol": r["rol"],
            "activo": "1" if r["activo"] else "0", "ultimo_acceso": r.get("ultimo_acceso")}


def usuario_actual(db=None):
    token = request.cookies.get(COOKIE)
    if not token:
        return None

    def buscar(d):
        r = d.fila("SELECT u.* FROM sesiones s JOIN usuarios u ON u.id = s.usuario_id "
                   "WHERE s.token_hash = ? AND s.expira > ? AND u.activo = 1", (_hash_token(token), time.time()))
        if r:
            d.ejecutar("UPDATE sesiones SET expira = ? WHERE token_hash = ?", (time.time() + DURACION_SESION, _hash_token(token)))
            return _usuario_publico(r)
        return None

    if db:
        return buscar(db)
    with Db() as d:
        return buscar(d)


def requiere_usuario(admin=False):
    u = usuario_actual()
    if not u:
        raise ErrorAPI(401, "Sesión vencida. Volvé a ingresar.")
    if admin and u["rol"] != "admin":
        raise ErrorAPI(403, "Solo un administrador puede hacer esto.")
    return u


def leer_json():
    if "application/json" not in (request.content_type or ""):
        raise ErrorAPI(415, "Se esperaba JSON.")
    datos = request.get_json(silent=True)
    if datos is None:
        raise ErrorAPI(400, "JSON inválido.")
    return datos


@app.before_request
def antes():
    if not request.path.startswith("/api/"):
        return None
    # Protección CSRF: los pedidos que modifican datos tienen que venir de esta misma página
    if request.method not in ("GET", "HEAD"):
        origen = request.headers.get("Origin")
        if origen and urllib.parse.urlparse(origen).netloc != request.host:
            return jsonify(ok=False, error="Origen no permitido."), 403
    if EN_VERCEL and not ES_PG:
        return jsonify(ok=False, error="Falta conectar la base de datos: en Vercel no existe la variable DATABASE_URL. "
                                       "Conectá Neon desde Storage y hacé Redeploy."), 503
    try:
        inicializar()
    except Exception as e:  # base caída o mal configurada
        app.logger.error("No se pudo conectar a la base (%s): %r", DB_VAR or "SQLite", e)
        return jsonify(ok=False, error=f"La base de datos rechazó la conexión ({type(e).__name__}). "
                                       "El detalle está en Vercel → Logs."), 503
    return None


@app.after_request
def despues(resp):
    resp.headers["Cache-Control"] = "no-store"
    resp.headers["X-Content-Type-Options"] = "nosniff"
    resp.headers["X-Frame-Options"] = "DENY"
    resp.headers["Referrer-Policy"] = "same-origin"
    resp.headers["X-Robots-Tag"] = "noindex, nofollow"
    return resp


@app.errorhandler(ErrorAPI)
def _err_api(e):
    return jsonify(ok=False, error=str(e)), e.codigo


@app.errorhandler(413)
def _err_grande(_):
    return jsonify(ok=False, error="Los datos superan el límite de 4 MB por envío (revisá las fotos cargadas)."), 413


@app.errorhandler(ValueError)
def _err_valor(e):
    return jsonify(ok=False, error=str(e)), 400


@app.errorhandler(Exception)
def _err_general(e):
    if hasattr(e, "code") and hasattr(e, "description"):  # errores HTTP de Flask (404, 405…)
        return jsonify(ok=False, error=e.description), e.code
    app.logger.exception("Error inesperado")
    return jsonify(ok=False, error="Error interno del servidor."), 500


# =============================================================================
# DATOS
# =============================================================================

def _es_activo(e):
    return str(e.get("activo", "1")).strip().lower() not in ("0", "false", "no", "inactivo", "baja")


def version_datos(db, bloquear=False):
    sql = "SELECT valor FROM meta WHERE clave = 'version_datos'" + (" FOR UPDATE" if bloquear and ES_PG else "")
    return int(db.fila(sql)["valor"])


def leer_datos(db):
    datos = {"empleados": [json.loads(r["datos"]) for r in db.filas("SELECT datos FROM empleados ORDER BY nombre")]}
    for r in db.filas("SELECT clave, datos FROM colecciones"):
        datos[r["clave"]] = json.loads(r["datos"])
    return datos


def guardar_respaldo(db, usuario, motivo):
    datos = leer_datos(db)
    if not datos["empleados"]:
        return
    db.ejecutar("INSERT INTO respaldos (fecha, usuario, motivo, empleados, datos) VALUES (?,?,?,?,?)",
                (ahora_iso(), usuario, motivo, len(datos["empleados"]), json.dumps(datos, ensure_ascii=False)))
    viejos = db.filas("SELECT id FROM respaldos ORDER BY id DESC")[RESPALDOS_A_CONSERVAR:]
    for r in viejos:
        db.ejecutar("DELETE FROM respaldos WHERE id = ?", (r["id"],))


def escribir_datos(db, datos):
    empleados = datos.get("empleados")
    if not isinstance(empleados, list):
        raise ValueError("Falta la lista de empleados.")
    ids = set()
    for e in empleados:
        if not isinstance(e, dict) or not e.get("id"):
            raise ValueError("Hay un empleado sin id.")
        if e["id"] in ids:
            raise ValueError(f"Id de empleado repetido: {e['id']}")
        ids.add(e["id"])
    ts = ahora_iso()
    existentes = {r["id"]: r["datos"] for r in db.filas("SELECT id, datos FROM empleados")}
    for e in empleados:
        js = json.dumps(e, ensure_ascii=False, sort_keys=True)
        if existentes.get(e["id"]) == js:
            continue
        db.ejecutar(
            "INSERT INTO empleados (id, legajo, nombre, planta, sector, activo, datos, actualizado) VALUES (?,?,?,?,?,?,?,?) "
            "ON CONFLICT (id) DO UPDATE SET legajo=excluded.legajo, nombre=excluded.nombre, planta=excluded.planta, "
            "sector=excluded.sector, activo=excluded.activo, datos=excluded.datos, actualizado=excluded.actualizado",
            (e["id"], str(e.get("legajo", "")), e.get("nombre", ""), e.get("planta", ""), e.get("sector", ""),
             1 if _es_activo(e) else 0, js, ts))
    borrados = set(existentes) - ids
    for emp_id in borrados:
        db.ejecutar("DELETE FROM empleados WHERE id = ?", (emp_id,))
    for clave in COLECCIONES:
        if clave in datos:
            db.ejecutar("INSERT INTO colecciones (clave, datos, actualizado) VALUES (?,?,?) "
                        "ON CONFLICT (clave) DO UPDATE SET datos=excluded.datos, actualizado=excluded.actualizado",
                        (clave, json.dumps(datos[clave], ensure_ascii=False), ts))
    nueva = version_datos(db) + 1
    db.ejecutar("UPDATE meta SET valor = ? WHERE clave = 'version_datos'", (str(nueva),))
    return nueva, len(borrados)


@app.get("/api/sesion")
def api_sesion():
    return jsonify(ok=True, usuario=usuario_actual(), version_app=APP_VERSION)


@app.post("/api/login")
def api_login():
    d = leer_json()
    user = str(d.get("username", "")).strip().lower()[:60]
    clave = str(d.get("password", ""))
    with Db() as db:
        if not db.fila("SELECT 1 AS x FROM usuarios"):
            raise ErrorAPI(503, "Todavía no hay usuarios. Configurá ADMIN_USERNAME y ADMIN_PASSWORD en Vercel y volvé a desplegar.")
        desde = (datetime.datetime.now(datetime.timezone.utc) - datetime.timedelta(minutes=15)).isoformat(timespec="seconds")
        fallidos = db.fila("SELECT COUNT(*) AS n FROM historial WHERE accion = 'login_fallido' AND usuario = ? AND fecha > ?",
                           (user, desde))["n"]
        if fallidos >= INTENTOS_LOGIN:
            raise ErrorAPI(429, "Demasiados intentos fallidos. Esperá 15 minutos y volvé a probar.")
        r = db.fila("SELECT * FROM usuarios WHERE username = ?", (user,))
        if not r or not r["activo"] or not verificar_password(clave, r["pass_hash"]):
            registrar(db, user, "login_fallido")
            ok = False
        else:
            token = secrets.token_urlsafe(32)
            db.ejecutar("INSERT INTO sesiones (token_hash, usuario_id, expira) VALUES (?,?,?)",
                        (_hash_token(token), r["id"], time.time() + DURACION_SESION))
            db.ejecutar("UPDATE usuarios SET ultimo_acceso = ? WHERE id = ?", (ahora_iso(), r["id"]))
            registrar(db, r["username"], "login")
            ok = True
    if not ok:
        time.sleep(0.6)
        raise ErrorAPI(401, "Usuario o contraseña incorrectos.")
    resp = jsonify(ok=True, usuario=_usuario_publico(r))
    seguro = request.headers.get("X-Forwarded-Proto", request.scheme) == "https"
    resp.set_cookie(COOKIE, token, httponly=True, samesite="Strict", secure=seguro, path="/")
    return resp


@app.post("/api/logout")
def api_logout():
    token = request.cookies.get(COOKIE)
    if token:
        with Db() as db:
            db.ejecutar("DELETE FROM sesiones WHERE token_hash = ?", (_hash_token(token),))
    resp = jsonify(ok=True)
    resp.delete_cookie(COOKIE, path="/")
    return resp


@app.get("/api/datos")
def api_leer():
    requiere_usuario()
    with Db() as db:
        return jsonify(ok=True, version=version_datos(db), datos=leer_datos(db))


@app.post("/api/datos")
def api_guardar():
    u = requiere_usuario()
    cuerpo = leer_json()
    with Db() as db:
        actual = version_datos(db, bloquear=True)
        base = cuerpo.get("version")
        if base is not None and int(base) != actual:
            return jsonify(ok=False, conflicto=True, version=actual,
                           error="Otra persona o ventana guardó cambios. Recargá la página para ver la última versión."), 409
        hoy = datetime.date.today().isoformat()
        if not db.fila("SELECT 1 AS x FROM respaldos WHERE fecha >= ? AND motivo = 'diario'", (hoy,)):
            guardar_respaldo(db, "sistema", "diario")
        nueva, borrados = escribir_datos(db, cuerpo.get("datos") or {})
        det = f"{len((cuerpo.get('datos') or {}).get('empleados', []))} empleados" + (f", {borrados} eliminados" if borrados else "")
        registrar(db, u["username"], "guardar", det)
    return jsonify(ok=True, version=nueva)


@app.post("/api/importar")
def api_importar():
    u = requiere_usuario(admin=True)
    datos = leer_json()
    if not isinstance(datos.get("empleados"), list):
        raise ValueError("El archivo no tiene una lista de empleados.")
    with Db() as db:
        version_datos(db, bloquear=True)
        guardar_respaldo(db, u["username"], "antes de restaurar")
        nueva, _ = escribir_datos(db, datos)
        registrar(db, u["username"], "importar", f"{len(datos['empleados'])} empleados")
    return jsonify(ok=True, version=nueva)


# =============================================================================
# USUARIOS
# =============================================================================

def _listar(db):
    return [_usuario_publico(r) for r in db.filas("SELECT * FROM usuarios ORDER BY nombre")]


@app.get("/api/usuarios")
def api_usuarios():
    requiere_usuario(admin=True)
    with Db() as db:
        return jsonify(ok=True, usuarios=_listar(db))


@app.post("/api/usuarios")
def api_guardar_usuario():
    actor = requiere_usuario(admin=True)
    d = leer_json()
    user = str(d.get("username") or "").strip().lower()
    nombre = str(d.get("nombre") or "").strip()
    clave = str(d.get("password") or "")
    rol = d.get("rol") if d.get("rol") in ("admin", "operador") else "operador"
    activo = 0 if str(d.get("activo", "1")) == "0" else 1
    uid = d.get("id")
    if not user or not nombre:
        raise ValueError("Usuario y nombre son obligatorios.")
    if not all(c.isalnum() or c in "._-" for c in user):
        raise ValueError("El usuario solo puede tener letras, números, punto, guion y guion bajo.")
    if clave and len(clave) < 8:
        raise ValueError("La contraseña debe tener al menos 8 caracteres.")
    with Db() as db:
        otro = db.fila("SELECT id FROM usuarios WHERE username = ?", (user,))
        if otro and otro["id"] != uid:
            raise ValueError("Ya existe un usuario con ese nombre.")
        if uid:
            if not db.fila("SELECT 1 AS x FROM usuarios WHERE id = ?", (uid,)):
                raise ValueError("Usuario inexistente.")
            if uid == actor["id"] and (rol != "admin" or not activo):
                raise ValueError("No podés quitarte el rol de administrador ni desactivarte.")
            db.ejecutar("UPDATE usuarios SET username=?, nombre=?, rol=?, activo=? WHERE id=?", (user, nombre, rol, activo, uid))
            if clave:
                db.ejecutar("UPDATE usuarios SET pass_hash=? WHERE id=?", (hash_password(clave), uid))
                if uid != actor["id"]:
                    db.ejecutar("DELETE FROM sesiones WHERE usuario_id=?", (uid,))
            if not activo:
                db.ejecutar("DELETE FROM sesiones WHERE usuario_id=?", (uid,))
            registrar(db, actor["username"], "editar_usuario", user)
        else:
            if not clave:
                raise ValueError("La contraseña es obligatoria para un usuario nuevo.")
            uid = "u" + secrets.token_hex(6)
            db.ejecutar("INSERT INTO usuarios (id, username, nombre, rol, activo, pass_hash, creado) VALUES (?,?,?,?,?,?,?)",
                        (uid, user, nombre, rol, activo, hash_password(clave), ahora_iso()))
            registrar(db, actor["username"], "crear_usuario", user)
        return jsonify(ok=True, id=uid, usuarios=_listar(db))


@app.delete("/api/usuarios/<uid>")
def api_eliminar_usuario(uid):
    actor = requiere_usuario(admin=True)
    if uid == actor["id"]:
        raise ValueError("No podés eliminar tu propio usuario.")
    with Db() as db:
        r = db.fila("SELECT username FROM usuarios WHERE id = ?", (uid,))
        if not r:
            raise ValueError("Usuario inexistente.")
        db.ejecutar("DELETE FROM sesiones WHERE usuario_id = ?", (uid,))
        db.ejecutar("DELETE FROM usuarios WHERE id = ?", (uid,))
        registrar(db, actor["username"], "eliminar_usuario", r["username"])
        return jsonify(ok=True, usuarios=_listar(db))


@app.post("/api/password")
def api_password():
    u = requiere_usuario()
    d = leer_json()
    nueva = str(d.get("nueva") or "")
    if len(nueva) < 8:
        raise ValueError("La contraseña nueva debe tener al menos 8 caracteres.")
    with Db() as db:
        r = db.fila("SELECT pass_hash FROM usuarios WHERE id = ?", (u["id"],))
        if not r or not verificar_password(str(d.get("actual") or ""), r["pass_hash"]):
            raise ValueError("La contraseña actual no es correcta.")
        db.ejecutar("UPDATE usuarios SET pass_hash = ? WHERE id = ?", (hash_password(nueva), u["id"]))
        registrar(db, u["username"], "cambiar_password")
    return jsonify(ok=True)


@app.get("/api/historial")
def api_historial():
    requiere_usuario(admin=True)
    with Db() as db:
        return jsonify(ok=True, historial=db.filas(
            "SELECT fecha, usuario, accion, detalle FROM historial ORDER BY id DESC LIMIT 300"))


@app.get("/api/respaldos")
def api_respaldos():
    requiere_usuario(admin=True)
    with Db() as db:
        return jsonify(ok=True, respaldos=db.filas(
            "SELECT id, fecha, usuario, motivo, empleados FROM respaldos ORDER BY id DESC"))


# =============================================================================
# DESCARGAS (Excel, CSV, backup)
# =============================================================================

CAMPOS = [
    ("legajo", "Legajo", 12), ("nombre", "Apellido y nombre", 34), ("dni", "DNI", 13), ("cuit", "CUIT", 16),
    ("empresa", "Empresa", 12), ("planta", "Planta", 12), ("sector", "Sector", 12), ("puesto", "Puesto", 22),
    ("categoria", "Categoría", 13), ("fechaNac", "Fecha nac.", 12), ("inicio", "Ingreso", 12), ("fin", "Egreso", 12),
    ("_antiguedad", "Antigüedad", 12), ("_estado", "Estado", 10), ("celular", "Celular personal", 16),
    ("celularLab", "Celular laboral", 16), ("cobertura", "Cobertura médica", 16), ("idioma", "Idioma", 16),
    ("direccion", "Dirección", 40), ("banco", "Banco", 14), ("cbu", "CBU", 24), ("alias", "Alias CBU", 18),
    ("consultora", "Consultora", 16), ("consultoraInicio", "Inicio consultora", 14), ("consultoraFin", "Fin consultora", 14),
]
FECHAS = {"fechaNac", "inicio", "fin", "consultoraInicio", "consultoraFin"}


def _fecha_ar(v):
    return f"{v[8:10]}/{v[5:7]}/{v[0:4]}" if isinstance(v, str) and len(v) == 10 and v[4] == "-" else (v or "")


def _antiguedad(e):
    try:
        ini = datetime.date.fromisoformat(e.get("inicio", ""))
    except ValueError:
        return ""
    fin = datetime.date.today()
    if not _es_activo(e) and e.get("fin"):
        try:
            fin = datetime.date.fromisoformat(e["fin"])
        except ValueError:
            pass
    m = (fin.year - ini.year) * 12 + fin.month - ini.month - (1 if fin.day < ini.day else 0)
    if m < 0:
        return ""
    a, m = divmod(m, 12)
    return f"{a}a {m}m" if a else f"{m}m"


def _fila(e):
    out = []
    for k, _, _ in CAMPOS:
        if k == "_antiguedad":
            out.append(_antiguedad(e))
        elif k == "_estado":
            out.append("Activo" if _es_activo(e) else "Baja")
        elif k in FECHAS:
            out.append(_fecha_ar(e.get(k, "")))
        else:
            v = e.get(k, "") or ""
            # Evita que Excel interprete un valor como fórmula
            out.append("'" + v if isinstance(v, str) and v[:1] in ("=", "+", "-", "@") else v)
    return out


def _empleados():
    with Db() as db:
        return leer_datos(db)["empleados"]


def _nombre(ext):
    return f"LOGITEx_Personal_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.{ext}"


@app.get("/api/exportar/csv")
def api_csv():
    requiere_usuario()
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow([c for _, c, _ in CAMPOS])
    for e in _empleados():
        w.writerow(_fila(e))
    return Response("﻿" + buf.getvalue(), mimetype="text/csv; charset=utf-8",
                    headers={"Content-Disposition": f'attachment; filename="{_nombre("csv")}"'})


@app.get("/api/exportar/excel")
def api_excel():
    requiere_usuario()
    if not XLSX:
        raise ErrorAPI(500, "Falta openpyxl en requirements.txt.")
    empleados = _empleados()
    libro = openpyxl.Workbook()
    h = libro.active
    h.title = "Personal"
    linea = Side(style="thin", color="C0C8D8")
    borde = Border(left=linea, right=linea, top=linea, bottom=linea)
    for col, (_, enc, ancho) in enumerate(CAMPOS, 1):
        c = h.cell(row=1, column=col, value=enc)
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0F3D73"), borde
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        h.column_dimensions[get_column_letter(col)].width = ancho
    h.row_dimensions[1].height = 30
    h.freeze_panes = "C2"
    par, baja = PatternFill("solid", fgColor="EEF2F8"), PatternFill("solid", fgColor="FDECEC")
    for n, e in enumerate(empleados, 2):
        relleno = baja if not _es_activo(e) else (par if n % 2 == 0 else None)
        for col, v in enumerate(_fila(e), 1):
            c = h.cell(row=n, column=col, value=v)
            c.border = borde
            if relleno:
                c.fill = relleno
    h.auto_filter.ref = f"A1:{get_column_letter(len(CAMPOS))}{max(1, len(empleados) + 1)}"
    r = libro.create_sheet("Resumen")
    r.column_dimensions["A"].width, r.column_dimensions["B"].width = 26, 14
    r["A1"] = "Resumen de personal — LOGITEx RRHH"
    r["A1"].font = Font(bold=True, size=14, color="0F3D73")
    activos = [e for e in empleados if _es_activo(e)]
    fila = 3
    for et, v in (("Total en nómina", len(empleados)), ("Activos", len(activos)), ("Bajas", len(empleados) - len(activos)),
                  ("Exportado", datetime.datetime.now().strftime("%d/%m/%Y %H:%M"))):
        r.cell(row=fila, column=1, value=et).font = Font(bold=True)
        r.cell(row=fila, column=2, value=v)
        fila += 1
    for titulo, clave in (("Activos por planta", "planta"), ("Activos por sector", "sector")):
        fila += 1
        r.cell(row=fila, column=1, value=titulo).font = Font(bold=True, color="0F3D73")
        fila += 1
        conteo = {}
        for e in activos:
            conteo[e.get(clave) or "(sin dato)"] = conteo.get(e.get(clave) or "(sin dato)", 0) + 1
        for k in sorted(conteo):
            r.cell(row=fila, column=1, value=k)
            r.cell(row=fila, column=2, value=conteo[k])
            fila += 1
    buf = io.BytesIO()
    libro.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=_nombre("xlsx"),
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/api/backup")
def api_backup():
    u = requiere_usuario()
    with Db() as db:
        datos = {"_version": APP_VERSION, "_exportadoEn": ahora_iso(), "_exportadoPor": u["username"], **leer_datos(db)}
        registrar(db, u["username"], "descargar_backup")
    nombre = f"LOGITEx_backup_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.json"
    return Response(json.dumps(datos, ensure_ascii=False, indent=1), mimetype="application/json",
                    headers={"Content-Disposition": f'attachment; filename="{nombre}"'})


@app.get("/api/respaldos/<int:rid>")
def api_respaldo(rid):
    requiere_usuario(admin=True)
    with Db() as db:
        r = db.fila("SELECT fecha, datos FROM respaldos WHERE id = ?", (rid,))
    if not r:
        raise ErrorAPI(404, "Respaldo inexistente.")
    nombre = f"LOGITEx_respaldo_{r['fecha'][:10]}_{rid}.json"
    return Response(r["datos"], mimetype="application/json",
                    headers={"Content-Disposition": f'attachment; filename="{nombre}"'})


# ── La interfaz ──
@app.get("/")
@app.get("/index.html")
def raiz():
    resp = send_file(os.path.join(DIR_RAIZ, "index.html"), mimetype="text/html")
    resp.headers["Strict-Transport-Security"] = "max-age=63072000"
    return resp


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=int(os.environ.get("PORT", 5000)), debug=False)
